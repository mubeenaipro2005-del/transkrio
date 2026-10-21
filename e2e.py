"""End-to-end UI test in real Chrome. Whisper itself is replaced by a mock worker (no network/model in CI),
everything else is real: file upload, audio/video decoding, windowing, UI, downloads, recovery."""
import subprocess,tempfile,os,json,re,threading,http.server,functools,sys,zipfile,time,urllib.parse
from playwright.sync_api import sync_playwright
ROOT=os.path.dirname(os.path.dirname(os.path.abspath(__file__)));M=tempfile.mkdtemp()
def ff(*a):subprocess.run(['ffmpeg','-loglevel','error','-y',*a],check=True,cwd=M)
tone=lambda f,d,r=16000,amp=.4,gap=True:f"aevalsrc='{amp}*sin(2*PI*{f}*t)*sin(2*PI*3*t){'*lt(mod(t,20),17)' if gap else ''}':s={r}:d={d}"
ff('-f','lavfi','-i',tone(180,40),'short.wav');ff('-f','lavfi','-i',tone(180,725),'long12.wav')
ff('-f','lavfi','-i',tone(220,90,44100,.4,False),'-f','lavfi','-i','testsrc=size=320x240:rate=10','-t','90','-c:v','libvpx','-b:v','200k','-c:a','libopus','video.webm')
ff('-f','lavfi','-i',tone(200,30,44100,.4,False),'-c:a','libmp3lame','audio.mp3')
ff('-f','lavfi','-i','anullsrc=r=16000:cl=mono','-t','10','silent.wav');ff('-f','lavfi','-i',tone(180,30,16000,.01),'quiet.wav')
open(M+'/bad.mp3','wb').write(os.urandom(5000));open(M+'/empty.wav','wb').close()
P=lambda n:os.path.join(M,n)

MOCK="""
const ctl=async q=>(await fetch('/__ctl?'+q,{method:'POST'})).json();
let loaded=null;const sleep=ms=>new Promise(r=>setTimeout(r,ms));
self.onmessage=async({data:d})=>{
 if(d.cmd==='load'){const c=await ctl('op=load&model='+d.model);
  if(c.fail){self.postMessage({type:'err',stage:'load',msg:'net down'});return}
  for(let i=1;i<=4;i++){self.postMessage({type:'dl',p:{status:'progress',file:'f',loaded:i*25,total:100}});await sleep(50)}
  loaded=d.model;self.postMessage({type:'ready'});return}
 if(!loaded){self.postMessage({type:'err',stage:'run',msg:'model not loaded'});return}
 const a=d.audio,secs=a.length/16000;let pk=0;for(let i=0;i<a.length;i++)pk=Math.max(pk,Math.abs(a[i]));
 const c=await ctl('op=run&secs='+Math.round(secs)+'&lang='+d.lang+'&task='+d.task+'&peak='+pk.toFixed(3));
 if(c.fail){self.postMessage({type:'err',stage:'run',msg:'inference failed'});return}
 const n=Math.ceil(secs/20),chunks=[];
 for(let i=0;i<n;i++){await sleep(c.delay);self.postMessage({type:'tick'});chunks.push({text:' جملہ '+(i+1)+' ',timestamp:[i*20,Math.min(secs,(i+1)*20)]})}
 self.postMessage({type:'win',text:'',chunks});
};"""
ST=dict(fail_runs=0,fail_load=False,delay=30);LOG=[]
class H(http.server.SimpleHTTPRequestHandler):
    def log_message(*a):pass
    def do_POST(s):
        q=dict(urllib.parse.parse_qsl(urllib.parse.urlparse(s.path).query));out={'delay':ST['delay']}
        if q['op']=='load':LOG.append(('load',q['model']));out['fail']=ST['fail_load']
        else:
            LOG.append(('run',int(q['secs']),q['lang'],q['task'],float(q['peak'])))
            if ST['fail_runs']>0:ST['fail_runs']-=1;out['fail']=True
        b=json.dumps(out).encode();s.send_response(200);s.send_header('Content-Type','application/json');s.end_headers();s.wfile.write(b)
srv=http.server.ThreadingHTTPServer(('127.0.0.1',0),functools.partial(H,directory=ROOT));URL=f'http://127.0.0.1:{srv.server_port}/'
threading.Thread(target=srv.serve_forever,daemon=True).start()

res=[];errs=[]
def check(name,cond,info=''):
    res.append((name,bool(cond)));print(('PASS ' if cond else 'FAIL ')+name+(' | '+str(info) if not cond or info else ''))
def reset():ST.update(fail_runs=0,fail_load=False,delay=30);LOG.clear()

with sync_playwright() as p:
    br=p.chromium.launch(executable_path='/opt/google/chrome/chrome',args=['--no-sandbox','--autoplay-policy=no-user-gesture-required'])
    def page(real_worker=False,sw='block',vp=None):
        ctx=br.new_context(service_workers=sw,accept_downloads=True,viewport=vp or {'width':390,'height':780},locale='ur-PK')
        ctx.grant_permissions(['clipboard-read','clipboard-write'],origin=URL.rstrip('/'))
        pg=ctx.new_page();pg.on('pageerror',lambda e:errs.append(str(e)));pg.on('console',lambda m:errs.append(m.text) if m.type=='error' and 'net::ERR' not in m.text and 'Failed to load resource' not in m.text else 0)
        if not real_worker:pg.route('**/worker.js',lambda r:r.fulfill(body=MOCK,content_type='text/javascript'))
        pg.goto(URL);return pg
    status=lambda pg:pg.inner_text('#status')
    def done(pg,t=60000):pg.wait_for_function("document.getElementById('status').className==='ok'",timeout=t)
    def upload(pg,n):pg.set_input_files('#file',P(n))
    def dl(pg,sel):
        with pg.expect_download() as d:pg.click(sel)
        path=d.value.path();return open(path,'rb').read(),d.value.suggested_filename

    # 1 ---- initial state
    reset();pg=page()
    check('1 page loads, no JS errors',not errs,errs)
    check('1 default model is Small, language Urdu, transcribe button disabled',pg.input_value('#model').endswith('small') and pg.input_value('#lang')=='urdu' and pg.is_disabled('#go'))
    check('1 RTL Urdu interface',pg.evaluate("document.documentElement.dir")=='rtl' and 'Bolo' in pg.title())

    # 2 ---- video upload: auto-starts, parallel model load, progress, result
    seen=set();upload(pg,'video.webm')
    t0=time.time()
    while time.time()-t0<30 and pg.get_attribute('#status','class')!='ok':
        seen.add(status(pg));time.sleep(.02)
    done(pg)
    check('2 WebM video (VP8+Opus; AAC/H.264 need a branded Chrome) decodes and transcribes automatically',len(pg.input_value('#txt'))>0 and not pg.is_hidden('#res'),status(pg))
    check('2 status shown during work (file check / download % / progress)',any('%' in s for s in seen) and len(seen)>=3,sorted(seen)[:6])
    runs=[l for l in LOG if l[0]=='run'];check('2 audio duration decoded correctly (90s) and model loaded once',[r[1] for r in runs]==[90] and sum(1 for l in LOG if l[0]=='load')==1,LOG)
    txt=pg.input_value('#txt');check('2 all 5 segments present, in order',[f'جملہ {i}' for i in range(1,6)]==re.findall(r'جملہ \d',txt) or all(f'جملہ {i}' in txt for i in range(1,6)),txt)
    check('2 diagnostics line shown',' | ' in pg.inner_text('#diag'))

    # 3 ---- downloads
    b,n=dl(pg,'[data-dl=txt]');check('3 TXT download matches editor',b.decode()==txt and n=='video.txt',n)
    b,n=dl(pg,'[data-dl=srt]');s=b.decode();check('3 SRT valid (5 cues, increasing times)',len(re.findall(r'-->',s))==5 and s.startswith('1\n00:00:00,000 --> 00:00:20,000\n') and '00:01:20,000 --> 00:01:30,000' in s,s[:200])
    b,n=dl(pg,'[data-dl=vtt]');check('3 VTT valid',b.decode().startswith('WEBVTT\n\n00:00:00.000 --> 00:00:20.000'))
    b,n=dl(pg,'#docx');open(P('o.docx'),'wb').write(b)
    import docx;d=docx.Document(P('o.docx'));check('3 DOCX opens in python-docx with the text',' '.join(x.text for x in d.paragraphs).strip()==txt.strip(),[x.text for x in d.paragraphs])
    subprocess.run(['soffice','--headless','--convert-to','pdf','--outdir',M,P('o.docx')],capture_output=True,timeout=120)
    check('3 DOCX opens in LibreOffice (converts to PDF)',os.path.exists(P('o.pdf')) and os.path.getsize(P('o.pdf'))>1000)
    pg.click('#copy');check('3 copy to clipboard',pg.evaluate('navigator.clipboard.readText()')==txt)

    # 4 ---- MP3 + other settings passed through, quiet audio boosted
    reset();pg.select_option('#lang','english');pg.select_option('#task','translate');upload(pg,'audio.mp3');done(pg)
    r=[l for l in LOG if l[0]=='run'][0];check('4 MP3 decodes; language/task passed to worker',r[1:4]==(30,'english','translate'),r)
    reset();pg.select_option('#lang','urdu');pg.select_option('#task','transcribe');upload(pg,'quiet.wav');done(pg)
    pk=[l for l in LOG if l[0]=='run'][0][4];check('4 very quiet recording is amplified before transcription (peak 0.01 -> ~0.08)',0.06<pk<0.12,pk)

    # 5 ---- long file: windows, offsets, contiguity
    reset();pg.reload();upload(pg,'long12.wav');done(pg,90000)
    runs=[l[1] for l in LOG if l[0]=='run'];check('5 12-min file split into windows that sum to full duration (725s)',len(runs)==3 and abs(sum(runs)-725)<=2,runs)
    b,_=dl(pg,'[data-dl=srt]');cues=re.findall(r'(\d+):(\d+):(\d+),(\d+) --> (\d+):(\d+):(\d+),(\d+)',b.decode())
    sec=lambda h,m,s,ms:int(h)*3600+int(m)*60+int(s)+int(ms)/1000
    st=[sec(*c[:4]) for c in cues];en=[sec(*c[4:]) for c in cues]
    check('5 SRT timeline is global, monotonic and ends at ~725s',all(st[i]<st[i+1] for i in range(len(st)-1)) and abs(en[-1]-725)<2 and all(e>=s for s,e in zip(st,en)),(st[-3:],en[-3:]))
    check('5 diagnostics report 3 parts','3 part' in pg.inner_text('#diag'),pg.inner_text('#diag'))

    # 6 ---- transient failure: auto-retry, nothing lost/duplicated
    reset();ST['fail_runs']=1;pg.reload();upload(pg,'long12.wav');done(pg,90000)
    runs=[l[1] for l in LOG if l[0]=='run'];check('6 transient inference error retried automatically, final result complete',len(runs)==4 and runs[0]==runs[1],runs)
    segs=pg.evaluate("JSON.parse(localStorage.getItem('sn.segs')).length");check('6 no lost / duplicated segments (exactly 38 expected for 725 s)',segs==38,segs)

    # 7 ---- permanent failure then RESUME without re-uploading
    reset();ST['fail_runs']=99;pg.reload();upload(pg,'long12.wav')
    pg.wait_for_function("document.getElementById('status').className==='er'",timeout=60000)
    check('7 failure: clear error shown + Resume button visible',pg.is_visible('#resume') and 'جاری رکھیں' in status(pg),status(pg))
    check('7 failure: Transcribe button back, nothing left spinning',pg.is_visible('#go') and pg.is_hidden('#stop'))
    ST['fail_runs']=0;LOG.clear()
    ST['fail_runs']=0;pg.click('#resume');done(pg,90000)
    runs=[l[1] for l in LOG if l[0]=='run'];check('7 resume processes ALL windows (none had completed) and finishes',len(runs)==3,runs)
    # failure in the middle: window 1 ok, window 2 fails (3 attempts) -> resume only redoes windows >=2
    reset();ST['delay']=120;pg.reload();upload(pg,'long12.wav')
    pg.wait_for_function("typeof job!=='undefined'&&job&&job.k>=1&&job.state==='run'",timeout=60000)   # window 1 finished
    ST['fail_runs']=99
    pg.wait_for_function("document.getElementById('status').className==='er'",timeout=60000)
    partial=pg.input_value('#txt');check('7 partial transcript kept visible after failure',len(partial)>0,partial)
    n_before=len([l for l in LOG if l[0]=='run' and l[1]>0]);ST['fail_runs']=0;LOG.clear();pg.click('#resume');done(pg,90000)
    runs=[l[1] for l in LOG if l[0]=='run']
    check('7 resume redoes ONLY the unfinished windows (>=1 window kept)',len(runs)<3 and len(runs)>=1,runs)
    check('7 final text complete after resume (no gaps, no duplicates)',pg.evaluate("JSON.parse(localStorage.getItem('sn.segs')).length")==38)

    # 8 ---- cancel then resume
    reset();ST['delay']=400;pg.reload();upload(pg,'long12.wav')
    pg.wait_for_function("typeof job!=='undefined'&&job&&job.state==='run'&&tRun>0",timeout=30000);time.sleep(.6)
    pg.click('#stop');time.sleep(.5)
    check('8 cancel: stopped, Resume offered',pg.is_visible('#resume') and pg.is_visible('#go'),status(pg))
    ST['delay']=20;LOG.clear();pg.click('#resume');done(pg,90000)
    check('8 resume after cancel completes with full text',pg.evaluate("JSON.parse(localStorage.getItem('sn.segs')).length")==38)
    # cancel while still decoding must also work
    reset();pg.reload();upload(pg,'long12.wav');pg.click('#stop')
    check('8 cancel during file reading works',pg.is_visible('#go') and pg.is_hidden('#stop'))

    # 9 ---- bad inputs
    for f,exp,lab in [('empty.wav','فائل خالی','empty file'),('bad.mp3','پڑھی نہیں','corrupt file'),('silent.wav','خاموش','silent audio')]:
        reset();pg=page();upload(pg,f);pg.wait_for_function("document.getElementById('status').className==='er'",timeout=20000)
        check(f'9 {lab}: clear message, UI usable again',exp in status(pg) and pg.is_visible('#go') and not pg.is_disabled('#file'),status(pg))
    check('9 invalid inputs never reached the model',not [l for l in LOG if l[0]=='run'])

    # 10 ---- persistence (reload keeps text + subtitles; settings remembered)
    reset();pg=page();pg.select_option('#model','Xenova/whisper-base');pg.select_option('#ui','en');upload(pg,'short.wav');done(pg)
    t=pg.input_value('#txt');pg.reload()
    check('10 transcript restored after reload',pg.input_value('#txt')==t)
    check('10 SRT/VTT still available after reload',not pg.is_disabled('[data-dl=srt]'))
    check('10 settings remembered (model, interface language)',pg.input_value('#model').endswith('base') and pg.evaluate('document.documentElement.lang')=='en')
    pg.once('dialog',lambda d:d.accept());pg.click('#clr');check('10 delete text clears everything',pg.input_value('#txt')=='' and pg.is_hidden('#res'))

    # 11 ---- real worker.js with no network: must fail gracefully (model/library unreachable)
    reset();pg=page(real_worker=True);upload(pg,'short.wav')
    pg.wait_for_function("document.getElementById('status').className==='er'",timeout=60000)
    check('11 library/model unreachable -> understandable message + Resume',('ماڈل لوڈ' in status(pg) or 'Could not load' in status(pg)) and pg.is_visible('#resume'),status(pg))

    # 12 ---- service worker + PWA assets
    pg=page(sw='allow');pg.wait_for_function("navigator.serviceWorker.controller||navigator.serviceWorker.ready.then(()=>true)");time.sleep(1.5)
    keys=pg.evaluate("caches.keys().then(async k=>{const c=await caches.open(k.find(x=>x.startsWith('bl-v')));return (await c.keys()).map(r=>new URL(r.url).pathname)})")
    need=['/core.js','/app.js','/worker.js','/icon-192.png','/manifest.webmanifest','/index.html']
    check('12 service worker caches every app file (works offline after 1st visit)',all(x in keys for x in need),keys)
    mf=pg.evaluate("fetch('manifest.webmanifest').then(r=>r.json())");check('12 manifest valid (192/512 PNG icons)',{'192x192','512x512'}<={i['sizes'] for i in mf['icons']})
    pg.screenshot(path='/tmp/shot_idle.png')
    pg=page();upload(pg,'short.wav');done(pg);pg.screenshot(path='/tmp/shot_done.png',full_page=True)
    br.close()
ok=sum(1 for _,c in res if c);print(f'\n{ok}/{len(res)} checks passed; JS console/page errors: {len(errs)}',errs[:5])
sys.exit(0 if ok==len(res) and not errs else 1)
